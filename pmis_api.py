"""
pmis_api.py — Explorateur de l'API PMIS (façon Postman)

- Catalogue des endpoints de l'API externe PMIS (Ed. 0 Rev. C) et de leurs filtres.
- Envoi d'une requête quelconque (statut, durée, taille, corps) sans lever d'erreur
  sur un code HTTP ≥ 400 : la réponse est affichée telle quelle, comme dans Postman.
- Mise à plat des réponses en tableaux lisibles (visites, mouvements, services).
- Génération d'une collection Postman v2.1 (sans aucun secret : URLs et identifiants
  sont des variables à renseigner dans Postman).
"""
from __future__ import annotations

import json
import time
from datetime import datetime

import pandas as pd
import requests

import pmis

# ═══════════════════════════════════════════════════════════════════════════════
#  CATALOGUE DES ENDPOINTS
# ═══════════════════════════════════════════════════════════════════════════════
VISIT_DATE_FIELDS = {"ETA": "eta", "ETD": "etd", "ATA": "ata", "ATD": "atd"}
SERVICE_NAMES = ["Pilotage", "Tug Assistance", "Mooring Service"]

ENDPOINTS = {
    "visits": {
        "method": "GET", "path": "/visits", "label": "Visites",
        "description": "Toutes les visites (escales) avec navire, mouvements et services "
                       "réservés. Au moins un couple de dates (après / avant) est requis.",
        "params": {
            "etaAfterDate": "ETA ≥ (AAAA-MM-JJ)", "etaBeforeDate": "ETA ≤",
            "etdAfterDate": "ETD ≥", "etdBeforeDate": "ETD ≤",
            "ataAfterDate": "ATA ≥", "ataBeforeDate": "ATA ≤",
            "atdAfterDate": "ATD ≥", "atdBeforeDate": "ATD ≤",
            "shipName": "Nom du navire (égalité)",
            "billable": "Facturable (true / false : billing_status vide)",
            "visitStatus": "Statut de la visite (STATUS_TYPE)",
            "page": "N° de page", "pageSize": "Éléments par page (max 50)",
        },
    },
    "services": {
        "method": "GET", "path": "/services", "label": "Services hors visite",
        "description": "Services réservés non rattachés à une visite.",
        "params": {
            "bookingAfterDate": "Réservation ≥ (AAAA-MM-JJ)",
            "bookingBeforeDate": "Réservation ≤",
            "shipName": "Nom du navire (égalité)", "serviceName": "Nom du service",
            "bookingStatus": "Statut de la réservation (STATUS_TYPE)",
            "page": "N° de page", "pageSize": "Éléments par page (max 50)",
        },
    },
    "billing": {
        "method": "PUT", "path": "/billing", "label": "Statut de facturation",
        "description": "⚠️ ÉCRITURE dans PMIS : met à jour le statut de facturation de "
                       "plusieurs visites (0 = non facturée, 1 = facturée).",
        "params": {},
    },
}


# ═══════════════════════════════════════════════════════════════════════════════
#  ENVOI D'UNE REQUÊTE
# ═══════════════════════════════════════════════════════════════════════════════
def clean_params(params: dict) -> dict:
    return {k: v for k, v in (params or {}).items() if v not in (None, "", [])}


def preview_url(client, endpoint: str, params: dict) -> str:
    base = client.backend_url or "https://PMIS_IP/pmisBackend/api"
    req = requests.Request("GET", base + ENDPOINTS[endpoint]["path"],
                           params=clean_params(params)).prepare()
    return req.url


def curl(client, endpoint: str, params: dict, body=None) -> str:
    """Commande cURL équivalente (le jeton est remplacé par $TOKEN)."""
    ep = ENDPOINTS[endpoint]
    lines = [f"curl -X {ep['method']} '{preview_url(client, endpoint, params)}'",
             '  -H "Authorization: Bearer $TOKEN"',
             "  -H 'Accept: application/json' -H 'Content-Type: application/json'"]
    if body is not None:
        lines.append(f"  -d '{json.dumps(body, ensure_ascii=False)}'")
    return " \\\n".join(lines)


def send(client, endpoint: str, params: dict | None = None, body=None) -> dict:
    """Envoie la requête et renvoie {status, ok, url, method, elapsed_ms, size, json,
    text, sent_at}. Lève pmis.PMISError seulement si le serveur est injoignable."""
    ep = ENDPOINTS[endpoint]
    if not client.access_token:
        client.login()
    url = client.backend_url + ep["path"]
    q = clean_params(params)

    def _do():
        return requests.request(ep["method"], url, params=q or None,
                                json=body if ep["method"] != "GET" else None,
                                headers=client._headers(), verify=client.verify_ssl,
                                timeout=client.timeout)

    t0 = time.perf_counter()
    try:
        r = _do()
        if r.status_code == 401:            # jeton expiré → refresh puis nouvel essai
            client.refresh()
            t0 = time.perf_counter()
            r = _do()
    except requests.RequestException as e:
        raise pmis.PMISError(f"Serveur PMIS injoignable ({url}) : {e}") from e
    elapsed = (time.perf_counter() - t0) * 1000
    try:
        payload = r.json()
    except ValueError:
        payload = None
    return {"status": r.status_code, "ok": r.ok, "url": r.url, "method": ep["method"],
            "endpoint": endpoint, "params": q, "body": body,
            "elapsed_ms": round(elapsed), "size": len(r.content or b""),
            "json": payload, "text": None if payload is not None else (r.text or "")[:5000],
            "sent_at": datetime.now().strftime("%H:%M:%S")}


def send_all_pages(client, endpoint: str, params: dict, max_pages: int = 40) -> dict:
    """Comme send(), mais parcourt toutes les pages et fusionne les lignes."""
    first = send(client, endpoint, {**params, "pageSize": 50, "page": None})
    if not first["ok"] or not isinstance(first["json"], dict):
        return first
    data = first["json"].get("data") or {}
    rows, seen = list(data.get("rows") or []), set()
    key = "visit_id" if endpoint == "visits" else "booking_id"
    seen.update(r.get(key) for r in rows)
    total = data.get("total") or 0
    page = data.get("page") if isinstance(data.get("page"), int) else 1
    pages, elapsed, size = 1, first["elapsed_ms"], first["size"]
    while rows and len(rows) < total and pages < max_pages:
        page += 1
        nxt = send(client, endpoint, {**params, "pageSize": 50, "page": page})
        nrows = ((nxt["json"] or {}).get("data") or {}).get("rows") or [] if nxt["ok"] else []
        new = [r for r in nrows if r.get(key) not in seen]
        if not new:
            break
        seen.update(r.get(key) for r in new)
        rows += new
        pages += 1
        elapsed += nxt["elapsed_ms"]
        size += nxt["size"]
    merged = dict(first)
    merged["json"] = {**first["json"], "data": {**data, "rows": rows, "page": "toutes",
                                                "pages_lues": pages}}
    merged.update(elapsed_ms=elapsed, size=size)
    return merged


# ═══════════════════════════════════════════════════════════════════════════════
#  RÉPONSES → TABLEAUX LISIBLES
# ═══════════════════════════════════════════════════════════════════════════════
def rows_of(resp: dict) -> list:
    js = resp.get("json")
    if isinstance(js, dict):
        data = js.get("data", js)
        if isinstance(data, dict):
            return data.get("rows") or []
        if isinstance(data, list):
            return data
    return js if isinstance(js, list) else []


def _dt(v) -> str:
    return pmis._dt_fr(v) if v else ""


def visits_table(rows: list[dict]) -> pd.DataFrame:
    out = []
    for v in rows:
        s = pmis.visit_summary(v)
        ship = v.get("ship") or {}
        p = pmis.visit_params(v)
        term, mode = pmis.resolve_terminal(p["legs"], None)
        codes = [((b.get("service_type") or {}).get("imo_code") or "")
                 for b in (v.get("resource_bookings") or [])]
        out.append({
            "Visit ID": s["visit_id"], "N° escale": s["business_id"], "Navire": s["navire"],
            "IMO": s["imo"], "Type": ((ship.get("ships_type") or {}).get("description") or ""),
            "GT": ship.get("gross_tonnage"), "LOA": ship.get("loa"),
            "Largeur": ship.get("width"), "Client": s["client"], "Activité": s["activite"],
            "Statut": s["statut"], "Facturable": s["billable"],
            "billing_status": v.get("billing_status"),
            "ETA": s["eta"], "ATA": s["ata"], "ETD": s["etd"], "ATD": s["atd"],
            "Postes": " → ".join(dict.fromkeys(l["location"] for l in p["legs"]
                                               if l.get("location"))),
            "Terminal": term if mode != "défaut" else "",
            "Mouvements": len(v.get("movements") or []),
            "Pilotage": codes.count("PILO"), "Remorquage": codes.count("TOWG"),
            "Lamanage": codes.count("MOOR"),
        })
    return pd.DataFrame(out)


def movements_table(rows: list[dict]) -> pd.DataFrame:
    out = []
    for v in rows:
        name = ((v.get("ship") or {}).get("target_name") or "").strip()
        for m in (v.get("movements") or []):
            out.append({
                "Visit ID": v.get("visit_id"), "Navire": name,
                "Mouvement ID": m.get("movement_id"),
                "Type": (m.get("movements_type") or {}).get("movement_type_name"),
                "De": (m.get("location_from") or {}).get("name") or "",
                "Vers": (m.get("location_to") or {}).get("name") or "",
                "ETM": _dt(m.get("etm")), "Début": _dt(m.get("start_dt")),
                "Passage": _dt(m.get("cross_dt")), "Fin": _dt(m.get("end_dt")),
                "Services": len(m.get("booking_ids") or []),
            })
    return pd.DataFrame(out)


def bookings_table(bookings: list[dict], visit=None) -> pd.DataFrame:
    out = []
    for b in bookings:
        st_ = b.get("service_type") or {}
        u = b.get("resource_usage_details") or {}
        res = b.get("resources") or []
        acts = {((r.get("resource_properties") or {}).get("usage_details") or {})
                .get("activity_type") for r in res} - {None, ""}
        names = []
        for r in res:
            d = r.get("resource_details") or {}
            names.append((d.get("target_name") or d.get("alias") or " ".join(
                x for x in (d.get("name"), d.get("surname")) if x)).strip())
        out.append({
            **({"Visit ID": visit.get("visit_id"),
                "Navire": ((visit.get("ship") or {}).get("target_name") or "").strip()}
               if visit else {}),
            "Booking ID": b.get("booking_id"), "Service": st_.get("custom_name"),
            "Code IMO": st_.get("imo_code"), "Mouvement ID": b.get("fk_movement_id"),
            "Réservé le": _dt(b.get("booking_dt")),
            "Début prévu": _dt(u.get("estimated_start")),
            "Début réel": _dt(u.get("actual_start")),
            "Fin réelle": _dt(u.get("actual_completion")),
            "Activité": ", ".join(sorted(acts)),
            "Ressources": ", ".join(n for n in names if n),
            "Annulé le": _dt(b.get("cancellation_dt")),
        })
    return pd.DataFrame(out)


def services_of_visits(rows: list[dict]) -> pd.DataFrame:
    frames = [bookings_table(v.get("resource_bookings") or [], v) for v in rows]
    frames = [f for f in frames if not f.empty]
    return pd.concat(frames, ignore_index=True) if frames else pd.DataFrame()


def search(df: pd.DataFrame, text: str) -> pd.DataFrame:
    if df.empty or not text:
        return df
    hay = df.astype(str).agg(" ".join, axis=1).str.lower()
    return df[hay.str.contains(text.strip().lower(), regex=False)]


# ═══════════════════════════════════════════════════════════════════════════════
#  COLLECTION POSTMAN v2.1 (sans secret)
# ═══════════════════════════════════════════════════════════════════════════════
def _pm_url(var: str, path: str, params: dict | None = None, active: dict | None = None):
    active = clean_params(active)
    query = [{"key": k, "value": str(active.get(k, "")), "disabled": k not in active,
              "description": desc} for k, desc in (params or {}).items()]
    raw = "{{" + var + "}}/" + path
    if active:
        raw += "?" + "&".join(f"{k}={v}" for k, v in active.items())
    return {"raw": raw, "host": ["{{" + var + "}}"], "path": [path], "query": query}


def postman_collection(current: dict | None = None) -> dict:
    """Collection Postman v2.1 : login (enregistre le jeton), refresh, visites,
    services, statut de facturation (+ la requête courante de l'explorateur).
    Aucun secret : URLs et identifiants sont des variables de collection vides."""
    login_test = [
        "const j = pm.response.json(); const d = j.data || j;",
        "pm.collectionVariables.set('access_token', d.accessToken || '');",
        "pm.collectionVariables.set('refresh_token', d.refreshToken || '');",
        "pm.test('Jeton reçu', () => pm.expect(d.accessToken).to.be.a('string'));",
    ]
    json_hdr = [{"key": "Content-Type", "value": "application/json"},
                {"key": "Accept", "value": "application/json"}]
    items = [
        {"name": "🔐 Auth — Login", "event": [{"listen": "test", "script": {
            "type": "text/javascript", "exec": login_test}}],
         "request": {"auth": {"type": "noauth"}, "method": "POST", "header": json_hdr,
                     "body": {"mode": "raw", "raw": json.dumps(
                         {"username": "{{username}}", "password": "{{password}}"}, indent=2)},
                     "url": {"raw": "{{auth_url}}/login", "host": ["{{auth_url}}"],
                             "path": ["login"]},
                     "description": "Connexion : enregistre accessToken / refreshToken dans "
                                    "les variables de la collection (utiliser « mail » à la "
                                    "place de « username » si nécessaire)."}},
        {"name": "🔐 Auth — Refresh", "event": [{"listen": "test", "script": {
            "type": "text/javascript", "exec": login_test}}],
         "request": {"auth": {"type": "noauth"}, "method": "POST", "header": json_hdr,
                     "body": {"mode": "raw", "raw": json.dumps(
                         {"refreshToken": "{{refresh_token}}"}, indent=2)},
                     "url": {"raw": "{{auth_url}}/refresh", "host": ["{{auth_url}}"],
                             "path": ["refresh"]}}},
    ]
    today = datetime.now().strftime("%Y-%m-%d")
    defaults = {"visits": {"etaAfterDate": today[:8] + "01", "etdBeforeDate": today,
                           "pageSize": 50},
                "services": {"bookingAfterDate": today[:8] + "01",
                             "bookingBeforeDate": today, "pageSize": 50}}
    for key in ("visits", "services"):
        ep = ENDPOINTS[key]
        items.append({"name": f"📄 {ep['label']} — GET {ep['path']}", "request": {
            "method": "GET", "header": json_hdr, "description": ep["description"],
            "url": _pm_url("backend_url", key, ep["params"], defaults[key])}})
    items.append({"name": "⚠️ Statut de facturation — PUT /billing", "request": {
        "method": "PUT", "header": json_hdr, "description": ENDPOINTS["billing"]["description"],
        "body": {"mode": "raw", "raw": json.dumps(
            [{"visit_id": 0, "billing_status": 1}], indent=2)},
        "url": {"raw": "{{backend_url}}/billing", "host": ["{{backend_url}}"],
                "path": ["billing"]}}})
    if current and current.get("endpoint") in ("visits", "services"):
        ep = ENDPOINTS[current["endpoint"]]
        items.append({"name": f"⭐ Requête de l'explorateur — {ep['label']}", "request": {
            "method": "GET", "header": json_hdr,
            "url": _pm_url("backend_url", current["endpoint"], ep["params"],
                           current.get("params"))}})
    return {
        "info": {"name": "NWM — PMIS External API",
                 "description": "Collection générée par le simulateur NWM. Renseignez les "
                                "variables auth_url, backend_url, username et password "
                                "(aucun secret n'est inclus), puis lancez « Auth — Login ».",
                 "schema": "https://schema.getpostman.com/json/collection/v2.1.0/"
                           "collection.json"},
        "auth": {"type": "bearer", "bearer": [{"key": "token", "value": "{{access_token}}",
                                               "type": "string"}]},
        "variable": [{"key": k, "value": v} for k, v in (
            ("auth_url", "https://PMIS_IP/pmisAuthServer/api"),
            ("backend_url", "https://PMIS_IP/pmisBackend/api"),
            ("username", ""), ("password", ""), ("access_token", ""),
            ("refresh_token", ""))],
        "item": items,
    }


def tables_excel(tables: dict) -> bytes:
    """Classeur Excel : une feuille par tableau (nom → DataFrame)."""
    import io
    from openpyxl.styles import Font, PatternFill
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        for name, df in tables.items():
            if df is None or df.empty:
                continue
            df.to_excel(xw, sheet_name=name[:31], index=False)
            ws = xw.sheets[name[:31]]
            for c in ws[1]:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor="0B3C5D")
            for col in ws.columns:
                w = max(len(str(c.value or "")) for c in col)
                ws.column_dimensions[col[0].column_letter].width = min(max(w + 2, 8), 40)
            ws.freeze_panes = "A2"
        if not xw.book.worksheets:
            pd.DataFrame({"Info": ["Aucune donnée"]}).to_excel(xw, sheet_name="Vide",
                                                               index=False)
    return buf.getvalue()
